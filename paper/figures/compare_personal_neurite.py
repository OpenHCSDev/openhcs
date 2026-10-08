"""Compare current native well measurements, without processing or tuning images.

The identity key is evaluation-only: never give it to a blind analysis author.
This entrypoint selects one run and an explicit well-aggregation protocol, never
a mixture of attempts. Its outputs do not establish segmentation accuracy.
Previously published comparison tables remain immutable historical evidence;
this reader consumes the current unit-bearing OpenHCS summary declaration.
"""

import argparse
from abc import ABC, abstractmethod
import csv
from dataclasses import asdict, dataclass, field, fields
import hashlib
import json
import math
from pathlib import Path
from statistics import mean, stdev
from typing import ClassVar


@dataclass(frozen=True)
class SummaryEndpoint:
    """One endpoint declaration supplies native decoding and reference identity."""

    label: str
    reference_column: str
    native_column: str
    unit: str = field(default="count", kw_only=True)

    def native_value(self, summary, cells):
        return float(summary[self.native_column])

    def reference_value(self, headers, values):
        return float(values[headers.index(self.reference_column)])

    def aggregate_values(self, values):
        return None if any(value is None for value in values) else mean(values)

    def treatment_effect(self, control_points, points):
        """Summarize a curve's own controls; undefined ratios remain undefined."""
        baseline = self.aggregate_values(control_points)
        value = self.aggregate_values(points)
        delta = None if baseline is None or value is None else value - baseline
        fold = None if baseline is None or value is None else RatioEndpoint.ratio(value, baseline)
        result = {"control_mean": baseline,
                  "control_sd": None if baseline is None else stdev(control_points),
                  "treatment_mean": value,
                  "treatment_sd": None if value is None else stdev(points),
                  "delta": delta, "fold_change": fold,
                  "fractional_change": None if fold is None else fold - 1}
        result["direction"] = ("undefined" if delta is None else
                               "increase" if delta > 0 else
                               "decrease" if delta < 0 else "unchanged")
        return result


class PerCellEndpoint(SummaryEndpoint):
    """An exported per-cell mean is undefined when the observed count is zero."""

    def native_value(self, summary, cells):
        return super().native_value(summary, cells) if cells else None

    def reference_value(self, headers, values):
        count = float(values[headers.index("Number of Cells (Neurite Outgrowth)")])
        return None if RatioEndpoint.ratio(0, count) is None else super().reference_value(headers, values)


class CellMeanEndpoint(PerCellEndpoint):
    """Mean of per-cell endpoints, not a pooled graph-segment statistic."""

    def native_value(self, summary, cells):
        values = [float(row[self.native_column]) for row in cells]
        if any(not math.isfinite(value) or value < 0 for value in values):
            raise ValueError(f"Invalid cell endpoint: {self.native_column}")
        return mean(values) if values else None


@dataclass(frozen=True)
class RatioEndpoint(SummaryEndpoint):
    """A declared ratio of exported totals, not a ratio of cell means."""

    reference_denominator: str
    native_denominator: str

    @staticmethod
    def ratio(numerator, denominator):
        if not math.isfinite(denominator) or denominator < 0:
            raise ValueError("Endpoint ratio requires a nonnegative finite denominator")
        if denominator == 0:
            return None  # Undefined, not a zero ratio or a rejected zero-growth field.
        return numerator / denominator

    def native_value(self, summary, cells):
        return self.ratio(super().native_value(summary, cells),
                          float(summary[self.native_denominator]))

    def reference_value(self, headers, values):
        return self.ratio(super().reference_value(headers, values),
                          float(values[headers.index(self.reference_denominator)]))


@dataclass(frozen=True)
class WellEndpoints:
    mean_outgrowth: float | None = field(metadata={"endpoint": PerCellEndpoint(
        "Mean outgrowth per cell / control", "Mean Outgrowth Per Cell (Neurite Outgrowth)", "mean_outgrowth_per_cell", unit="micrometers/cell")})
    cell_count: float = field(metadata={"endpoint": SummaryEndpoint(
        "Detected cells / control", "Number of Cells (Neurite Outgrowth)", "number_of_cells")})
    total_outgrowth: float = field(metadata={"endpoint": SummaryEndpoint(
        "Total outgrowth / control", "Total Outgrowth (Neurite Outgrowth)", "total_outgrowth", unit="micrometers")})
    branches_per_cell: float | None = field(metadata={"endpoint": PerCellEndpoint(
        "Branches per cell / control", "Mean Branches Per Cell (Neurite Outgrowth)", "mean_branches_per_cell", unit="branches/cell")})
    mean_process_length: float | None = field(metadata={"endpoint": CellMeanEndpoint(
        "Mean cell process length / control", "Cell: Mean Process Length (Neurite Outgrowth)", "mean_process_length", unit="micrometers")})
    median_process_length: float | None = field(metadata={"endpoint": CellMeanEndpoint(
        "Mean cell median process length / control", "Cell: Median Process Length (Neurite Outgrowth)", "median_process_length", unit="micrometers")})
    total_branches: float = field(metadata={"endpoint": SummaryEndpoint(
        "Total branches / control", "Total Branches (Neurite Outgrowth)", "total_branches")})
    total_processes: float = field(metadata={"endpoint": SummaryEndpoint(
        "Total primary processes / control", "Total Processes (Neurite Outgrowth)", "total_processes")})
    branches_per_process: float | None = field(metadata={"endpoint": RatioEndpoint(
        "Branches per primary process / control", "Total Branches (Neurite Outgrowth)", "total_branches",
        "Total Processes (Neurite Outgrowth)", "total_processes", unit="branches/primary_process")})

    def paired_values(self, guided, metrics):
        """Absolute endpoints and guided-minus/over-blind, never accuracy scores."""
        result = {}
        for metric in metrics:
            blind_value, guided_value = getattr(self, metric), getattr(guided, metric)
            defined = blind_value is not None and guided_value is not None
            result.update({f"openhcs_{metric}": blind_value,
                           f"guided_openhcs_{metric}": guided_value,
                           f"guided_minus_blind_{metric}": guided_value - blind_value if defined else None,
                           f"guided_over_blind_{metric}": RatioEndpoint.ratio(guided_value, blind_value) if defined else None})
        return result


METRICS = tuple(field.name for field in fields(WellEndpoints))
DEFAULT_METRICS = ("mean_outgrowth", "cell_count")


def endpoint_declarations():
    return {item.name: item.metadata["endpoint"] for item in fields(WellEndpoints)}


def read_rows(path):
    with path.open(newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames is None:
            raise ValueError(f"Missing CSV header: {path}")
        return list(reader)


def write_rows(path, rows):
    with path.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=tuple(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


@dataclass(frozen=True)
class NativeSummary:
    """The measured endpoint projection of one exported native plane."""

    coordinate_unit: ClassVar[str] = "micrometers"
    well: str
    site: str | None
    endpoints: WellEndpoints

    @staticmethod
    def cells_path(path):
        return path.with_name(path.name.replace("_neurite_outgrowth_summary_", "_neurite_outgrowth_cells_"))

    @classmethod
    def read(cls, path):
        rows = read_rows(path)
        if len(rows) != 1:
            raise ValueError(f"Expected one native plane summary: {path}")
        row = rows[0]
        if (int(row["neurite_channel_index"]), int(row["cell_body_channel_index"]),
                int(row["nuclear_channel_index"])) != (1, 1, 0):
            raise ValueError(f"Unexpected channel assignment: {path}")
        if row["coordinate_unit"] != cls.coordinate_unit:
            raise ValueError(f"Expected calibrated micrometer measurements: {path}")
        count, length = int(row["number_of_cells"]), float(row["total_outgrowth"])
        measured_mean = float(row["mean_outgrowth_per_cell"])
        if count < 0 or any(not math.isfinite(value) or value < 0 for value in (length, measured_mean)):
            raise ValueError(f"Invalid cell count or length measurement: {path}")
        if not math.isclose(length / count if count else 0, measured_mean, rel_tol=1e-9):
            raise ValueError(f"Inconsistent mean outgrowth: {path}")
        if (int(row["z_index"]), int(row["timepoint"])) != (1, 1):
            raise ValueError(f"Unexpected acquisition plane: {path}")
        cells = read_rows(cls.cells_path(path))
        if len(cells) != count or len({int(cell["cell"]) for cell in cells}) != count:
            raise ValueError(f"Cell identities disagree with summary: {path}")
        if any(cell["coordinate_unit"] != cls.coordinate_unit or
               cell["well"] != row["well"] or cell.get("site") != row.get("site")
               for cell in cells):
            raise ValueError(f"Cell measurements use another source or calibration: {path}")
        for cell_column, summary_column in (("total_outgrowth", "total_outgrowth"),
                                            ("branches", "total_branches"),
                                            ("processes", "total_processes")):
            cell_values = [float(cell[cell_column]) for cell in cells]
            if any(not math.isfinite(value) or value < 0 for value in cell_values):
                raise ValueError(f"Invalid cell {cell_column}: {path}")
            summary_value = float(row[summary_column])
            if ((count == 0 and summary_value != 0) or
                    not math.isclose(sum(cell_values), summary_value, rel_tol=1e-9, abs_tol=1e-9)):
                raise ValueError(f"Cell {cell_column} does not reconcile to summary: {path}")
        if not math.isclose(float(row["total_branches"]) / count if count else 0,
                            float(row["mean_branches_per_cell"]), rel_tol=1e-9):
            raise ValueError(f"Inconsistent branch denominator: {path}")
        values = {name: declaration.native_value(row, cells)
                  for name, declaration in endpoint_declarations().items()}
        if not all(math.isfinite(value) and value >= 0
                   for value in values.values() if value is not None):
            raise ValueError(f"Invalid native endpoint: {path}")
        return cls(row["well"], row.get("site"), WellEndpoints(**values))


class WellAggregation(ABC):
    """Each sampling protocol owns its well endpoint and coverage rule."""

    name: str
    description: str
    figure_label: str
    process_length_description: str
    total_outgrowth_description: str

    @abstractmethod
    def aggregate(self, rows):
        """Return well endpoints without silently changing sampled area."""

    def load_planes(self, summaries):
        """Decode once, retaining the native well/site identity for pairing."""
        planes, paths = {}, []
        summary_paths = []
        for directory in self.summary_directories(summaries):
            found = sorted(directory.glob("*_neurite_outgrowth_summary_*details.csv"))
            if not found:
                raise ValueError(f"No measured native summaries in declared directory: {directory}")
            summary_paths.extend(found)
        for path in summary_paths:
            row = NativeSummary.read(path)
            identity = (row.well, row.site)
            if identity in planes:
                raise ValueError(f"Duplicate native well/site: {identity}")
            planes[identity] = row
            paths.append(path)
            paths.append(NativeSummary.cells_path(path))
        return planes, paths

    @staticmethod
    def summary_directories(summaries):
        """A single native directory or an explicitly declared directory sequence."""
        return (summaries,) if isinstance(summaries, Path) else tuple(summaries)

    def aggregate_planes(self, planes):
        grouped = {}
        for row in planes.values():
            grouped.setdefault(row.well, []).append(row)
        return {well: self.aggregate(rows) for well, rows in grouped.items()}

    def load(self, summaries):
        planes, paths = self.load_planes(summaries)
        return self.aggregate_planes(planes), paths


class MosaicWellAggregation(WellAggregation):
    name = "mosaic"
    description = "OpenHCS mosaic total length / detected cells; MetaXpress well export"
    figure_label = "Stitched-mosaic well means from a fixed-pipeline transfer evaluation"
    process_length_description = "Mean of per-cell mean/median root-partition lengths in the mosaic; not pooled segment statistics. MetaXpress cell-summary aggregation is not declared."
    total_outgrowth_description = "Total traced length in the site-collapsed mosaic, not a sum of overlapping field totals."

    def aggregate(self, rows):
        if len(rows) != 1 or rows[0].site is not None:
            raise ValueError("Mosaic protocol requires exactly one site-collapsed summary per well")
        row = rows[0]
        return row.endpoints


class SiteMeanWellAggregation(WellAggregation):
    name = "site-mean"
    description = "Unweighted mean of nine native site endpoints; MetaXpress well export"
    figure_label = "Nine-field well means from a fixed-pipeline transfer evaluation"
    process_length_description = "Mean of per-cell mean/median root-partition lengths within each site, then unweighted site mean; not pooled segment statistics. MetaXpress cell-summary aggregation is not declared."
    total_outgrowth_description = "Unweighted mean of field totals; overlap is not deduplicated and field totals are not summed as unique-well length."

    def aggregate(self, rows):
        if len(rows) != 9 or {row.site for row in rows} != {str(site) for site in range(1, 10)}:
            raise ValueError("Site-mean protocol requires nine unique measured sites, 1–9, for each well")
        return WellEndpoints(**{
            metric: declaration.aggregate_values([getattr(row.endpoints, metric) for row in rows])
            for metric, declaration in endpoint_declarations().items()
        })


def reference_rows(reference, workbook, metrics):
    """Use the existing well metadata; decode original workbook endpoints once."""
    rows = read_rows(reference)
    if workbook is not None:
        from openpyxl import load_workbook
        book = load_workbook(workbook, read_only=True, data_only=True)
        try:
            wanted = {(row["excel_sheet"], int(row["excel_row"])): row for row in rows}
            declarations = endpoint_declarations()
            for sheet in book:
                headers = None
                for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    if "Number of Cells (Neurite Outgrowth)" in values:
                        headers = values
                    target = wanted.get((sheet.title, number))
                    if target is None:
                        continue
                    if headers is None or values[0] != target["well"]:
                        raise ValueError(f"Workbook source row identity changed: {sheet.title}/{number}")
                    for metric in metrics:
                        target[metric] = declarations[metric].reference_value(headers, values)
                    del wanted[(sheet.title, number)]
            if wanted:
                raise ValueError("Reference metadata names missing workbook rows")
        finally:
            book.close()
    for row in rows:
        for metric in metrics:
            # Only the workbook endpoint decoder can establish an observed
            # zero denominator. Missing/blank CSV values are not observations.
            if workbook is not None and row[metric] is None:
                continue
            if row.get(metric) in (None, "") or not math.isfinite(float(row[metric])) or float(row[metric]) < 0:
                raise ValueError(f"Missing or invalid reference endpoint {metric}: {row['plate']}/{row['well']}")
    return rows


def compare(reference, key, summaries, output, aggregation, coded_plate, pipeline,
            *, metrics=None, workbook=None,
            guided_summaries=None, guided_pipeline=None, guided_label=None):
    guided_options = (guided_summaries, guided_pipeline, guided_label)
    if any(value is not None for value in guided_options) and not all(guided_options):
        raise ValueError("Guided evaluation requires summaries, pipeline and an explicit label")
    if guided_label is not None and not guided_label.strip():
        raise ValueError("Guided evaluation requires an explicit label")
    paired = guided_summaries is not None
    metrics = tuple(metrics) if metrics is not None else METRICS if paired else DEFAULT_METRICS
    inputs = [reference, key, pipeline]
    if workbook is not None:
        inputs.append(workbook)
    mapping = {}
    for image in json.loads(key.read_text())["images"]:
        plate, filename = image["coded_relative_path"].split("/")
        coded_well = filename.split("_")[0]
        # Original acquisition directory, not the randomized coded plate,
        # carries physical plate identity.
        physical_plate = Path(image["source_path"]).parents[1].name.rsplit("_", 1)[1]
        identity = (physical_plate, image["source_well"])
        coded = (plate, coded_well)
        if coded in mapping and mapping[coded] != identity:
            raise ValueError(f"Inconsistent source identity: {coded}")
        mapping[coded] = identity
    planes, paths = aggregation.load_planes(summaries)
    measured = aggregation.aggregate_planes(planes)
    native = {mapping[(coded_plate, well)]: asdict(endpoints) for well, endpoints in measured.items()}
    if len(native) != len(measured):
        raise ValueError("Multiple coded wells map to the same physical well")
    inputs.extend(paths)
    if paired:
        guided_planes, paths = aggregation.load_planes(guided_summaries)
        if not planes or planes.keys() != guided_planes.keys():
            raise ValueError("Blind and guided native well/site coverage must match exactly")
        guided_measured = aggregation.aggregate_planes(guided_planes)
        guided_native = {mapping[(coded_plate, well)]: asdict(endpoints)
                         for well, endpoints in guided_measured.items()}
        inputs.extend([guided_pipeline, *paths])
    joined = []
    for row in reference_rows(reference, workbook, metrics):
        identity = (row["plate"], row["well"])
        if identity not in native:
            continue
        joined.append({**row, **{f"openhcs_{metric}": native[identity][metric]
                                for metric in metrics}})
        if paired:
            joined[-1].update({"guided_label": guided_label,
                              **WellEndpoints(**native[identity]).paired_values(
                                  WellEndpoints(**guided_native[identity]), metrics)})
    if paired and (len(joined) != len(native) or
                   {(row["plate"], row["well"]) for row in joined} != native.keys()):
        raise ValueError("Reference must cover every paired physical well exactly once")
    methods = [("metaxpress", ""), ("openhcs", "openhcs_")]
    if paired:
        methods.append(("guided_openhcs", "guided_openhcs_"))
    effects = []
    for plate, condition in dict.fromkeys((row["plate"], row["condition"]) for row in joined):
        curve = [row for row in joined if (row["plate"], row["condition"]) == (plate, condition)]
        control = [row for row in curve if float(row["nominal_dose_uM"]) == 0]
        if len(control) != 2:
            raise ValueError(f"Missing curve-specific controls: {condition}")
        for dose in (0, 5, 10, 20, 40):
            drug = [row for row in curve if float(row["nominal_dose_uM"]) == dose]
            if len(drug) != 2:
                raise ValueError(f"Missing technical replicate: {condition}/{dose}")
            for metric in metrics:
                result = {"plate": drug[0]["plate"], "condition": condition,
                          "dose_uM": dose, "metric": metric,
                          "baseline_wells": ";".join(row["well"] for row in control),
                          "treatment_wells": ";".join(row["well"] for row in drug),
                          "n_control": len(control), "n_treatment": len(drug)}
                for method, prefix in methods:
                    column = f"{prefix}{metric}"
                    control_points = [None if row[column] is None else float(row[column]) for row in control]
                    points = [None if row[column] is None else float(row[column]) for row in drug]
                    effect = endpoint_declarations()[metric].treatment_effect(control_points, points)
                    result.update({f"{method}_{name}": value for name, value in effect.items()
                                   if paired or name != "direction"})
                native_change, reference_change = result["openhcs_fractional_change"], result["metaxpress_fractional_change"]
                result["fractional_change_difference"] = (None if native_change is None or reference_change is None
                                                           else native_change - reference_change)
                if paired:
                    result["guided_label"] = guided_label
                effects.append(result)
    hashes = {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs}
    output.mkdir(parents=True, exist_ok=True)
    write_rows(output / "joined_wells.csv", joined)
    write_rows(output / "treatment_effects.csv", effects)
    if paired:
        paired_sites = []
        for (well, site), plane in planes.items():
            physical_plate, physical_well = mapping[(coded_plate, well)]
            guided_plane = guided_planes[(well, site)]
            result = {"plate": physical_plate, "well": physical_well,
                      "coded_plate": coded_plate, "coded_well": well,
                      "site": site, "coordinate_unit": NativeSummary.coordinate_unit,
                      "guided_label": guided_label}
            result.update(plane.endpoints.paired_values(guided_plane.endpoints, metrics))
            paired_sites.append(result)
        write_rows(output / "paired_sites.csv", paired_sites)
    evidence = {
        "sources_sha256": hashes, "matched_wells": len(joined),
        "comparison": "fixed-run treatment evaluation; not a manual accuracy score",
        "aggregation_protocol": aggregation.name,
        "aggregation": aggregation.description,
        "coded_plate": coded_plate,
        "metrics": list(metrics),
        "process_length_aggregation": aggregation.process_length_description,
        "total_outgrowth_aggregation": aggregation.total_outgrowth_description,
        "branch_total_aggregation": "Total branches and total primary processes use the selected well aggregation; site-mean means field totals, not deduplicated whole-well counts.",
        "branch_process_ratio_aggregation": "OpenHCS: declared total branches / total primary processes in each native plane, then selected well aggregation (unweighted site mean for site-mean). MetaXpress: ratio of exported well totals; its internal site weighting and primary-process definition are unspecified. These are response proxies, not identical endpoint definitions.",
        "coordinate_unit": NativeSummary.coordinate_unit,
        "pipeline_source": str(pipeline),
        "protocol_figure_label": aggregation.figure_label,
        "generator_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "baseline": "each drug curve's own zero-dose DMSO wells",
    }
    directories = aggregation.summary_directories(summaries)
    if len(directories) > 1:
        evidence["native_summary_directories"] = [str(directory) for directory in directories]
    if paired:
        guided_directories = aggregation.summary_directories(guided_summaries)
        evidence["endpoint_units"] = {metric: endpoint_declarations()[metric].unit for metric in metrics}
        evidence["guided_openhcs"] = {"label": guided_label, "pipeline_source": str(guided_pipeline),
                                     "summaries": (str(guided_directories[0]) if len(guided_directories) == 1 else
                                                   [str(directory) for directory in guided_directories]),
                                     "matched_sites": len(planes),
                                     "role": "scientist-guided comparator, not ground truth",
                                     "aggregation_protocol": aggregation.name}
    (output / "source_evidence.json").write_text(json.dumps(evidence, indent=2) + "\n")
    print(f"Compared {len(joined)} physical wells; wrote {len(effects)} endpoint/dose rows to {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("reference", "key", "output", "pipeline"):
        parser.add_argument(f"--{name}", type=Path, required=True)
    parser.add_argument("--summaries", type=Path, nargs="+", required=True,
                        help="Declared native result directories; each well/site must occur exactly once across all.")
    parser.add_argument("--coded-plate", required=True,
                        help="Source plate identity in the evaluation key, not the output directory name")
    parser.add_argument("--metrics", nargs="+", choices=METRICS,
                        help="Defaults to mean outgrowth/count; guided comparisons default to all declared endpoints.")
    parser.add_argument("--reference-workbook", type=Path,
                        help="Original workbook supplies selected endpoints at the CSV's retained Excel row identities.")
    parser.add_argument("--guided-summaries", type=Path, nargs="+",
                        help="Optional scientist-guided native summaries; must match the blind well/site coverage and aggregation.")
    parser.add_argument("--guided-pipeline", type=Path)
    parser.add_argument("--guided-label", help="Explicit scientist-guided comparator label; never ground truth.")
    protocols = {protocol.name: protocol for protocol in WellAggregation.__subclasses__()}
    parser.add_argument("--aggregation", choices=protocols, required=True)
    args = parser.parse_args()
    compare(args.reference, args.key, args.summaries, args.output,
            protocols[args.aggregation](), args.coded_plate, args.pipeline,
            metrics=args.metrics, workbook=args.reference_workbook,
            guided_summaries=args.guided_summaries, guided_pipeline=args.guided_pipeline,
            guided_label=args.guided_label)
